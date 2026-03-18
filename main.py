#!/usr/bin/env python3
import asyncio, ipaddress, sys, csv, os

# ===== GUI DETECTION =====
GUI_AVAILABLE = True
try:
    from PyQt6.QtWidgets import (
        QApplication, QWidget, QVBoxLayout, QPushButton, QLineEdit,
        QTextEdit, QLabel, QMessageBox, QCheckBox, QFileDialog
    )
except:
    GUI_AVAILABLE = False

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table as PDFTable
from docx import Document

# ===== CONFIG =====
TIMEOUT = 1
MAX_CONCURRENT = 500
ULTRA_CONCURRENT = 1500  # mode rapide

SERVICE_PORTS = {
    21:"FTP",22:"SSH",23:"Telnet",25:"SMTP",53:"DNS",
    80:"HTTP",110:"POP3",143:"IMAP",443:"HTTPS",
    3306:"MySQL",3389:"RDP",5432:"PostgreSQL",
    6379:"Redis",8080:"HTTP Alt",8443:"HTTPS Alt"
}

def detect_service(port):
    return SERVICE_PORTS.get(port,"Unknown")

# ===== UTILS =====
def expand_targets(target):
    try:
        net = ipaddress.ip_network(target, strict=False)
        return [str(ip) for ip in net.hosts()]
    except:
        return [target]

def parse_ports(text):
    ports = set()
    for part in text.split(","):
        if "-" in part:
            start,end = part.split("-")
            ports.update(range(int(start), int(end)+1))
        else:
            ports.add(int(part.strip()))
    return sorted(ports)

# ===== SCAN =====
async def scan_port(ip, port, sem):
    async with sem:
        try:
            reader, writer = await asyncio.wait_for(
                asyncio.open_connection(ip, port), timeout=TIMEOUT
            )
            writer.close()
            await writer.wait_closed()
            return port
        except:
            return None

async def run_scan(target, ports, ultra=False):
    ips = expand_targets(target)
    concurrency = ULTRA_CONCURRENT if ultra else MAX_CONCURRENT
    sem = asyncio.Semaphore(concurrency)

    tasks = []
    for ip in ips:
        for port in ports:
            tasks.append(scan_port(ip, port, sem))

    results = await asyncio.gather(*tasks)

    final = []
    i = 0
    for ip in ips:
        for port in ports:
            if results[i]:
                final.append({
                    "ip": ip,
                    "port": port,
                    "service": detect_service(port)
                })
            i += 1

    return final

# ===== EXPORT =====
def export_docx(results, folder):
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, "scan_report.docx")

    doc = Document()
    doc.add_heading("Network Scan Report", 0)
    doc.add_paragraph(f"Total results: {len(results)}\n")

    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "IP"
    hdr[1].text = "Port"
    hdr[2].text = "Service"

    for r in results:
        row = table.add_row().cells
        row[0].text = r["ip"]
        row[1].text = str(r["port"])
        row[2].text = r["service"]

    doc.save(path)
    return path

def export_csv(results):
    with open("scan.csv","w",newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["ip","port","service"])
        writer.writeheader()
        writer.writerows(results)

# ===== CLI =====
def run_cli():
    print("\n=== Network Audit Tool ===\n")

    target = input("Target IP / CIDR: ")
    ports = parse_ports(input("Ports: "))
    ultra = input("Ultra fast mode? (y/n): ").lower() == "y"
    folder = input("DOCX output folder: ")

    print("\nScanning...\n")
    results = asyncio.run(run_scan(target, ports, ultra))

    if not results:
        print("No open ports found.")
    else:
        print("\n=== RESULTS ===\n")
        for r in results:
            print(f"[+] {r['ip']}:{r['port']} → {r['service']}")

        path = export_docx(results, folder)
        export_csv(results)

        print(f"\nSaved DOCX: {path}")
        print("Saved CSV: scan.csv")

# ===== GUI =====
if GUI_AVAILABLE:

    class ScannerGUI(QWidget):
        def __init__(self):
            super().__init__()
            self.setWindowTitle("Network Audit Tool")
            layout = QVBoxLayout()

            self.target = QLineEdit()
            self.target.setPlaceholderText("Target IP / CIDR")

            self.ports = QLineEdit("20-1024")

            self.ultra = QCheckBox("Ultra Fast Mode")

            self.folder_btn = QPushButton("Select DOCX Folder")
            self.folder_btn.clicked.connect(self.select_folder)
            self.folder_path = ""

            self.output = QTextEdit()

            btn = QPushButton("Start Scan")
            btn.clicked.connect(self.start_scan)

            layout.addWidget(QLabel("Target:"))
            layout.addWidget(self.target)

            layout.addWidget(QLabel("Ports:"))
            layout.addWidget(self.ports)

            layout.addWidget(self.ultra)
            layout.addWidget(self.folder_btn)
            layout.addWidget(btn)
            layout.addWidget(self.output)

            self.setLayout(layout)

        def select_folder(self):
            self.folder_path = QFileDialog.getExistingDirectory(self, "Select Folder")

        def start_scan(self):
            target = self.target.text()
            ports = parse_ports(self.ports.text())
            ultra = self.ultra.isChecked()

            self.output.clear()
            results = asyncio.run(run_scan(target, ports, ultra))

            if not results:
                self.output.append("No open ports found.")
                return

            self.output.append("=== RESULTS ===\n")

            for r in results:
                self.output.append(f"[+] {r['ip']}:{r['port']} → {r['service']}")

            folder = self.folder_path if self.folder_path else "."
            path = export_docx(results, folder)
            export_csv(results)

            self.output.append(f"\nDOCX saved: {path}")
            self.output.append("CSV saved: scan.csv")

# ===== ENTRY =====
if __name__ == "__main__":
    if GUI_AVAILABLE:
        try:
            app = QApplication(sys.argv)
            win = ScannerGUI()
            win.show()
            sys.exit(app.exec())
        except:
            run_cli()
    else:
        run_cli()
